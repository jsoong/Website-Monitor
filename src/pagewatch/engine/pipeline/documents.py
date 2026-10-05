"""Documents to HTML: PDF, DOCX and XLSX (spec: Fetch layer, document row).

Conversion is deterministic (same bytes, same HTML) because the viewer converts the stored raw
blob again to inject its highlights at the same DOM paths the extractor saw. Everything is
escaped; no markup from the document survives.

* PDF: text per page, one paragraph per text box, in pdfminer's reading order.
* DOCX: paragraphs and tables in document order (headings keep their level). Headers, footers,
  footnotes and text boxes are not read.
* XLSX: one table per visible sheet, capped at 5,000 rows.
"""

from __future__ import annotations

import io
from collections.abc import Callable
from datetime import date, datetime, time
from html import escape
from typing import Any

MAX_XLSX_ROWS = 5000
MAX_PDF_PAGES = 500
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # legacy .doc / .xls

Warn = Callable[[str], None]


class DocumentError(ValueError):
    """The bytes are not a readable document of the expected kind (a failed check)."""


def _shell(body: str) -> str:
    return f"<!DOCTYPE html><html><body>{body}</body></html>"


# -- PDF --------------------------------------------------------------------------------


def pdf_to_html(data: bytes, warn: Warn | None = None) -> str:
    warn = warn or (lambda _m: None)
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LAParams, LTTextContainer
    from pdfminer.pdfparser import PDFSyntaxError

    if not data.lstrip()[:5] == b"%PDF-":
        raise DocumentError("not a PDF (missing %PDF header)")
    parts: list[str] = []
    pages = 0
    try:
        for number, page in enumerate(
            extract_pages(io.BytesIO(data), laparams=LAParams(), maxpages=MAX_PDF_PAGES), start=1
        ):
            pages = number
            paragraphs: list[str] = []
            for element in page:
                if isinstance(element, LTTextContainer):
                    text = " ".join(
                        line.strip() for line in element.get_text().splitlines() if line.strip()
                    )
                    if text:
                        paragraphs.append(f"<p>{escape(text)}</p>")
            if paragraphs:
                inner = "".join(paragraphs)
                parts.append(f'<div class="pw-page" data-page="{number}">{inner}</div>')
    except DocumentError:
        raise
    except PDFSyntaxError as exc:
        raise DocumentError(f"unreadable PDF: {exc}") from exc
    except Exception as exc:  # pdfminer raises many types for damaged or encrypted files
        raise DocumentError(f"unreadable PDF: {type(exc).__name__}: {exc}"[:300]) from exc
    if pages >= MAX_PDF_PAGES:
        warn(f"PDF read up to {MAX_PDF_PAGES} pages")
    if not parts:
        warn("the PDF has no extractable text (a scan?)")
    return _shell("".join(parts))


# -- DOCX -------------------------------------------------------------------------------


def _heading_level(style: str) -> int | None:
    name = style.strip().lower()
    if name == "title":
        return 1
    if name.startswith("heading "):
        tail = name[8:].strip()
        if tail.isdigit():
            return min(max(int(tail), 1), 6)
    return None


def docx_to_html(data: bytes, warn: Warn | None = None) -> str:
    if data[:8] == _OLE_MAGIC:
        raise DocumentError("legacy .doc files are not supported; save the file as .docx")
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError(f"unreadable DOCX: {type(exc).__name__}: {exc}"[:300]) from exc
    out: list[str] = []
    in_list = False

    def close_list() -> None:
        nonlocal in_list
        if in_list:
            out.append("</ul>")
            in_list = False

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = " ".join(para.text.split())
            if not text:
                continue
            style = para.style.name if para.style is not None and para.style.name else ""
            level = _heading_level(style)
            if level is not None:
                close_list()
                out.append(f"<h{level}>{escape(text)}</h{level}>")
            elif style.lower().startswith("list"):
                if not in_list:
                    out.append("<ul>")
                    in_list = True
                out.append(f"<li>{escape(text)}</li>")
            else:
                close_list()
                out.append(f"<p>{escape(text)}</p>")
        elif tag == "tbl":
            close_list()
            table = Table(child, doc)
            rows: list[str] = []
            for row in table.rows:
                cells: list[str] = []
                last = None
                for cell in row.cells:
                    if cell._tc is last:  # merged cells repeat the same underlying cell
                        continue
                    last = cell._tc
                    cells.append(f"<td>{escape(' '.join(cell.text.split()))}</td>")
                rows.append(f"<tr>{''.join(cells)}</tr>")
            out.append(f"<table>{''.join(rows)}</table>")
    close_list()
    return _shell("".join(out))


# -- XLSX -------------------------------------------------------------------------------


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    return " ".join(str(value).split())


def xlsx_to_html(data: bytes, warn: Warn | None = None) -> str:
    warn = warn or (lambda _m: None)
    if data[:8] == _OLE_MAGIC:
        raise DocumentError("legacy .xls files are not supported; save the file as .xlsx")
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentError(f"unreadable XLSX: {type(exc).__name__}: {exc}"[:300]) from exc
    out: list[str] = []
    try:
        for ws in wb.worksheets:
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            out.append(f"<h2>{escape(ws.title)}</h2>")
            rows: list[str] = []
            capped = False
            for values in ws.iter_rows(values_only=True):
                cells = [_cell_text(v) for v in values]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells:
                    continue
                if len(rows) >= MAX_XLSX_ROWS:
                    capped = True
                    break
                rows.append("<tr>" + "".join(f"<td>{escape(c)}</td>" for c in cells) + "</tr>")
            if capped:
                warn(f"sheet {ws.title!r} read up to {MAX_XLSX_ROWS} rows")
            out.append(f"<table>{''.join(rows)}</table>")
    finally:
        wb.close()
    return _shell("".join(out))
