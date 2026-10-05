"""Builders for the document fixtures: a minimal PDF writer, DOCX and XLSX, and PNG pages.

The PDF writer emits real PDF 1.4 (Helvetica text at fixed positions, a cross-reference table),
so pdfminer reads exactly what a download would give it, without a PDF-writing dependency.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from typing import Any

from PIL import Image, ImageDraw


def _pdf_text(text: str) -> bytes:
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return safe.encode("latin-1", errors="replace")


def make_pdf(pages: Sequence[Sequence[str]]) -> bytes:
    """One page per entry; each string becomes its own paragraph (a text block 48 pt apart)."""
    objs: dict[int, bytes] = {}
    kids: list[int] = []
    next_id = 4  # 1 catalog, 2 page tree, 3 font
    for paragraphs in pages:
        stream = b"BT /F1 12 Tf\n"
        y = 740
        for text in paragraphs:
            stream += b"1 0 0 1 72 %d Tm (" % y + _pdf_text(text) + b") Tj\n"
            y -= 48
        stream += b"ET"
        content_id, page_id = next_id, next_id + 1
        next_id += 2
        objs[content_id] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        objs[page_id] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content_id
        )
        kids.append(page_id)
    objs[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objs[2] = (
        b"<< /Type /Pages /Kids ["
        + b" ".join(b"%d 0 R" % k for k in kids)
        + b"] /Count %d >>" % len(kids)
    )
    objs[3] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for num in sorted(objs):
        offsets[num] = out.tell()
        out.write(b"%d 0 obj\n" % num + objs[num] + b"\nendobj\n")
    xref = out.tell()
    size = max(objs) + 1
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % size)
    for num in range(1, size):
        out.write(b"%010d 00000 n \n" % offsets[num])
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (size, xref))
    return out.getvalue()


def make_docx(
    blocks: Sequence[tuple[str, Any]],
) -> bytes:
    """``("h1", "Title")``, ``("p", "text")``, ``("li", "item")``, ``("table", [[...], ...])``."""
    from docx import Document

    doc = Document()
    for kind, value in blocks:
        if kind.startswith("h"):
            doc.add_heading(value, level=int(kind[1:]))
        elif kind == "p":
            doc.add_paragraph(value)
        elif kind == "li":
            doc.add_paragraph(value, style="List Bullet")
        elif kind == "table":
            rows = value
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = str(cell)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_xlsx(sheets: dict[str, Sequence[Sequence[Any]]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def png_bytes(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def page_png(
    boxes: Sequence[tuple[int, int, int, int, str]] = (),
    size: tuple[int, int] = (1366, 900),
    background: str = "white",
) -> bytes:
    """A white page with filled rectangles ``(x, y, w, h, colour)``: a stand-in for a screenshot."""
    img = Image.new("RGB", size, background)
    draw = ImageDraw.Draw(img)
    for x, y, w, h, colour in boxes:
        draw.rectangle([x, y, x + w - 1, y + h - 1], fill=colour)
    return png_bytes(img)
